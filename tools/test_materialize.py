# -*- coding: utf-8 -*-
# SPDX-License-Identifier: MIT
# SPDX-FileCopyrightText: 2023-2026 x1shang
"""
frocloud · DIYPCmarket —— test_materialize.py

针对「合并单元格标签必须物化」这条规则的回归测试。
（2026-10-08 的 bug：gpu-brands.csv 33 行里 25 行 tier 为空 —— 合并段
 B3:B6=一线 只落在锚点 B3 上，其余格读到 None；同时跳过分区表头行的逻辑
 也漏了，导致后面的行整体错位一格。）

思路：不依赖真实数据，在**系统临时目录**里现造一个带合并段 + 分区表头 + 空锚点
合并段的小工作簿（绝不碰仓库里的 frocloud.xlsx），跑 export_csv.export_brands()，断言：
    ① 有值的合并段 → 该段每一行都拿到标签
    ② 空锚点的合并段 → 该段每一行如实为空（不瞎填）
    ③ 分区表头行不产出数据行（不把后面的行对齐顶偏）
    ④ 两个分区各自的联盟分组正确

用法：python tools/test_materialize.py     # 失败退出码 1
"""
from __future__ import annotations

import os
import sys
import tempfile

import openpyxl
from openpyxl.utils import column_index_from_string

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import export_csv as ex  # noqa: E402


def build_fixture(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "显卡品牌"
    ws["A2"], ws["B2"], ws["C2"] = "N", "品牌分级", "显卡品牌"

    # 分区 N：一线(B3:B6) / 次一线(B7:B8)，各段只在锚点写标签
    for r, (tier, brand) in enumerate(
            [("一线", "甲"), ("一线", "乙"), ("一线", "丙"), ("一线", "丁"),
             ("次一线", "戊"), ("次一线", "己")], start=3):
        if tier:
            ws.cell(row=r, column=2, value=tier)
        ws.cell(row=r, column=3, value=brand)
    ws.merge_cells("B3:B6")
    ws.merge_cells("B7:B8")

    # 分区 A：自带一行表头（这是必须被排除的行），然后 二线(B11:B13)
    ws["A9"], ws["B9"], ws["C9"] = "A", "品牌分级", "显卡品牌"
    for r, brand in enumerate(["庚", "辛", "壬"], start=11):
        ws.cell(row=r, column=3, value=brand)
    ws.cell(row=11, column=2, value="二线")
    ws.merge_cells("B11:B13")

    # 分区 I：整段合并但锚点是空的 —— 真空白，必须如实为空
    ws["A15"], ws["C15"] = "I", "癸"
    ws["C16"] = "子"
    ws.merge_cells("B15:B16")

    wb.save(path)


def main():
    tmp = os.path.join(tempfile.gettempdir(), "frocloud-fixture.xlsx")
    build_fixture(tmp)

    src = ex.Src(tmp)
    rep = ex.Report()
    header, rows = ex.export_brands(src, rep)

    print("导出结果：")
    for r in rows:
        print("   %-14s %-8s %s" % (r[0], r[1] or "(空)", r[2]))
    print()

    errors = []

    def expect(cond, msg):
        print(("   [OK] " if cond else "   [!!] ") + msg)
        if not cond:
            errors.append(msg)

    print("断言：① 有值合并段必须物化到每一行")
    for brand, want in [("甲", "一线"), ("乙", "一线"), ("丙", "一线"), ("丁", "一线"),
                        ("戊", "次一线"), ("己", "次一线"), ("庚", "二线"),
                        ("辛", "二线"), ("壬", "二线")]:
        got = next((r[1] for r in rows if r[2] == brand), None)
        expect(got == want, "%s → tier=%r（期望 %r）" % (brand, got, want))

    print("断言：② 空锚点的合并段必须如实为空")
    for brand in ("癸", "子"):
        got = next((r[1] for r in rows if r[2] == brand), None)
        expect(got == "", "%s → tier=%r（期望空字符串，不能瞎填）" % (brand, got))

    print("断言：③ 分区表头行不产出数据行")
    expect(all(r[2] != "显卡品牌" for r in rows), "没有 '显卡品牌' 数据行")

    print("断言：④ 联盟分组/行数正确")
    expect(len(rows) == 11, "数据行数 = %d（期望 11）" % len(rows))
    expect(next(r[0] for r in rows if r[2] == "庚") == "AMD 阵营", "A 分区归到 AMD 阵营")
    expect(next(r[0] for r in rows if r[2] == "癸") == "Intel 阵营", "I 分区归到 Intel 阵营")

    print()
    print("断言：⑤ 检查脚本也会拦住这种回归")
    print("   （下面这条 [ERR ] 是**故意**触发的：把 fixture 的 tier 清空，")
    print("     验证 check_brands 确实会报错。它写在 stderr 上，不计入本测试的失败。）")
    try:
        import check_csv as ck
        tables = {"gpu-brands.csv": [[r[0], r[1], r[2]] + r[3:] for r in rows]}
        # 故意把 tier 清空，模拟"没物化"的旧行为
        broken = [list(r) for r in tables["gpu-brands.csv"]]
        for r in broken:
            if r[2] in ("乙", "丙", "丁"):
                r[1] = ""
        ck.errors.clear(), ck.warnings.clear(), ck.notes.clear()
        ck.XLSX = tmp
        ck.check_brands({"gpu-brands.csv": broken}, False, label="【预期报错 · 回归测试用】")
        print("   （上面那条 [ERR ] 到此为止，属于预期输出）")
        expect(len(ck.errors) > 0, "把 tier 清空后 check_brands 报错（拦住绿着错）")
    except Exception as e:
        expect(False, "调用 check_brands 失败：%s" % e)

    print()
    if errors:
        print("结论：%d 条不通过" % len(errors))
        return 1
    print("结论：全部通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
