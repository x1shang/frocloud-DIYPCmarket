# -*- coding: utf-8 -*-
"""
frocloud · DIYPCmarket —— check_csv.py

校验 data/*.csv（月更流程的 CI：跑不过就返回非零退出码）。
对应 todo.txt 的清单：
    ① 列名集合一致
    ② date 可解析、(同一数据序列内) 单调不重复
    ③ 价格必须是数字
    ④ 无空行（字段数不符 / 全空行）
    ⑤ 表数 = 7（7 张数据表各一个 CSV）
再加几条这份数据特有的：
    ⑥ utf-8-sig 带 BOM（Windows 双击不乱码）
    ⑦ 天梯图 TOTAL = TIME SPY + Port Royal + TIME SPY(2) + Steel Nomad
    ⑧ 内存/存储的 AVG / MAX / MIN 与同组原始单价一致（用 xlsx 缓存值交叉验证）
    ⑨ 文本不以 = + - @ 开头（防公式注入）
    ⑩ 显卡行情「涨幅」与 CSV 重算值对照，并显式指出原公式基准列与"表内历史最低价"不一致的地方

用法：
    python tools/check_csv.py          # 校验，失败退出码 1
    python tools/check_csv.py -v       # 额外打印明细
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import os
import re
import sys
from collections import defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA = os.path.join(ROOT, "data")
CANDIDATES = [os.path.join(ROOT, "frocloud.xlsx"),
              os.path.join(ROOT, "frocloud-DIYPCmarket", "frocloud.xlsx")]
XLSX = next((p for p in CANDIDATES if os.path.exists(p)), CANDIDATES[0])

ENCODING = "utf-8-sig"
PRICE_HEADER = ["date", "category", "model", "spec", "price_cny", "channel", "note"]
TIER_HEADER = ["tier", "vendor", "platform", "model", "cores", "threads",
               "score_cpumark", "score_single", "note"]
GPU_TIER_HEADER = ["tier", "vendor", "series", "model", "score_timespy", "score_portroyal",
                   "score_timespy_2", "score_steelnomad", "score_total"]
BRAND_HEADER = ["vendor_group", "tier", "brand", "入门(丐版)", "低端", "中端",
                "高端", "旗舰", "超旗舰"]

EXPECTED = {
    "gpu-tier.csv": GPU_TIER_HEADER,
    "cpu-tier.csv": TIER_HEADER,
    "laptop-cpu-tier.csv": TIER_HEADER,
    "gpu-prices.csv": PRICE_HEADER,
    "cpu-prices.csv": PRICE_HEADER,
    "memory-prices.csv": PRICE_HEADER,
    "gpu-brands.csv": BRAND_HEADER,
}
PRICE_FILES = ["gpu-prices.csv", "cpu-prices.csv", "memory-prices.csv"]
TIER_FILES = ["gpu-tier.csv", "cpu-tier.csv", "laptop-cpu-tier.csv"]

DATE_MONTH = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
DATE_HALF = re.compile(r"^\d{4}-(0[1-9]|1[0-2])-下旬$")
FORMULA_RISK = re.compile(r"^[=+\-@]")
NUM = re.compile(r"^-?\d+(\.\d+)?$")

errors = []
warnings = []
notes = []


def err(msg, stream=None):
    """默认写到 stderr —— 这样被别的脚本（如 test_materialize.py）调用时，
    故意触发的报错不会污染调用方的 stdout。"""
    errors.append(msg)
    target = stream or sys.stderr
    try:
        print("  [ERR ] " + msg, file=target)
        target.flush()
    except Exception:
        # stderr 不可用时兜底到 stdout，绝不能因为打日志失败而吞掉错误
        print("  [ERR ] " + msg)


def warn(msg):
    warnings.append(msg)
    print("  [WARN] " + msg)


def ok(msg, verbose=True):
    if verbose:
        print("  [ OK ] " + msg)


def read_csv(path):
    with open(path, "rb") as f:
        raw = f.read()
    has_bom = raw.startswith(b"\xef\xbb\xbf")
    text = raw.decode(ENCODING)
    rows = list(csv.reader(text.splitlines()))
    return has_bom, rows


def norm(s):
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s)).strip()


def parse_date(s):
    if DATE_MONTH.match(s):
        y, m = s.split("-")
        return dt.date(int(y), int(m), 1)
    if DATE_HALF.match(s):
        y, m, _ = s.split("-")
        return dt.date(int(y), int(m), 28)   # 下旬按月末近似，仅用于排序
    return None


# ------------------------------------------------------------------ ①③④⑥⑨
def check_shape(verbose):
    print("\n1) 结构：文件数 / 编码 / 列名 / 空行 / 数字  [① ③ ④ ⑥ ⑨]")
    if not os.path.isdir(DATA):
        err("找不到 data/ 目录")
        return {}
    got = sorted(f for f in os.listdir(DATA) if f.endswith(".csv"))
    missing = sorted(set(EXPECTED) - set(got))
    extra = sorted(set(got) - set(EXPECTED))
    if missing:
        err("缺少 CSV：%s" % missing)
    if extra:
        err("多出未登记的 CSV：%s（请同步 EXPECTED 与 README）" % extra)
    ok("表数 = %d（期望 %d）" % (len(got), len(EXPECTED)))

    tables = {}
    for name in sorted(set(EXPECTED) & set(got)):
        path = os.path.join(DATA, name)
        has_bom, rows = read_csv(path)
        if not has_bom:
            err("%s 不是 utf-8-sig（缺 BOM），Windows 上双击 Excel 会乱码" % name)
        if not rows:
            err("%s 是空文件" % name)
            continue
        header = rows[0]
        if header != EXPECTED[name]:
            err("%s 列名不符\n         实际 %s\n         期望 %s" % (name, header, EXPECTED[name]))
        body = rows[1:]
        tables[name] = body
        nblank = sum(1 for r in body if not any(x.strip() for x in r))
        if nblank:
            err("%s 有 %d 个空行" % (name, nblank))
        nbad = [i + 2 for i, r in enumerate(body) if len(r) != len(header)]
        if nbad:
            err("%s 有字段数不符的行：第 %s 行" % (name, nbad[:5]))
        nrisk = 0
        for r in body:
            for x in r:
                if x and FORMULA_RISK.match(x):
                    nrisk += 1
        if nrisk:
            err("%s 有 %d 个以 = + - @ 开头的文本，存在公式注入风险" % (name, nrisk))
        # 数字列
        numcol = {"price_cny": "price", "cores": "int", "threads": "int",
                  "score_cpumark": "num", "score_single": "num", "score_timespy": "num",
                  "score_portroyal": "num", "score_timespy_2": "num",
                  "score_steelnomad": "num", "score_total": "num"}
        for col in header:
            if col not in numcol:
                continue
            idx = header.index(col)
            for i, r in enumerate(body):
                v = r[idx]
                if v == "":
                    continue
                if not NUM.match(v):
                    err("%s 第 %d 行 %s=%r 不是数字" % (name, i + 2, col, v))
                    break
                if col == "price_cny" and float(v) <= 0:
                    err("%s 第 %d 行 price_cny=%r 非正数" % (name, i + 2, v))
                    break
        ok("%-22s %4d 数据行，列 %d，BOM ✓" % (name, len(body), len(header)), verbose)
    return tables


# ------------------------------------------------------------------ ②
def check_dates(tables, verbose):
    print("\n2) 日期：可解析 / 单调 / 不重复  [②]")
    for name in PRICE_FILES:
        body = tables.get(name) or []
        if not body:
            err("%s 没有数据行" % name)
            continue
        h = EXPECTED[name]
        idx = h.index("date")
        # 唯一键 = 除 date / price_cny / note 外的所有列（内存/存储靠 spec 区分容量）
        keys = [c for c in h if c not in ("date", "price_cny", "note")]
        seen = defaultdict(list)
        bad = 0
        for i, r in enumerate(body):
            d = parse_date(r[idx])
            if d is None:
                err("%s 第 %d 行 date=%r 无法解析" % (name, i + 2, r[idx]))
                bad += 1
                continue
            seen[tuple(r[h.index(k)] for k in keys)].append(d)
        if bad:
            continue
        unsorted = [k for k, v in seen.items() if v != sorted(v)]
        if unsorted:
            err("%s 有 %d 组数据日期非单调（如 %s）" % (name, len(unsorted), unsorted[0]))
        dup = [(k, len(v)) for k, v in seen.items() if len(v) != len(set(v))]
        if dup:
            err("%s 有重复 (型号, 规格, date)：%s" % (name, dup[:3]))
        months = sorted({d for v in seen.values() for d in v})
        span = "%s ~ %s" % (months[0].strftime("%Y-%m"), months[-1].strftime("%Y-%m"))
        ok("%-22s %d 个数据序列 / %d 个月份，%s" % (name, len(seen), len(months), span), verbose)


# ------------------------------------------------------------------ ⑦
def check_tier_math(tables, verbose):
    print("\n3) 天梯图：TOTAL = TIME SPY + Port Royal + TIME SPY(2) + Steel Nomad  [⑦]")
    body = tables.get("gpu-tier.csv") or []
    h = GPU_TIER_HEADER
    bad = 0
    for i, r in enumerate(body):
        try:
            parts = [float(r[h.index(k)]) for k in
                     ("score_timespy", "score_portroyal", "score_timespy_2", "score_steelnomad")]
            total = float(r[h.index("score_total")])
        except (ValueError, IndexError):
            continue
        if abs(sum(parts) - total) > 0.5:
            bad += 1
            if bad <= 5:
                err("gpu-tier.csv 第 %d 行 %s：分量合计 %.0f ≠ TOTAL %.0f"
                    % (i + 2, r[h.index("model")], sum(parts), total))
    if not bad:
        ok("%d 行 TOTAL 全部等于四项之和" % len(body))
    # 核心 ≤ 线程
    for name in ("cpu-tier.csv", "laptop-cpu-tier.csv"):
        b = tables.get(name) or []
        hh = TIER_HEADER
        bad = 0
        for i, r in enumerate(b):
            try:
                c, t = int(r[hh.index("cores")]), int(r[hh.index("threads")])
            except (ValueError, IndexError):
                continue
            if c > t:
                bad += 1
                if bad <= 3:
                    err("%s 第 %d 行 %s：核心 %d > 线程 %d" % (name, i + 2, r[3], c, t))
        if not bad:
            ok("%s 核心/线程数自洽" % name)


# ------------------------------------------------------------------ ⑪
def check_brands(tables, verbose, label=""):
    """显卡品牌：vendor_group / tier 必须物化到每一行。

    原表里 B 列「品牌分级」是合并单元格（B3:B6=一线 …）—— 合并段的空白是
    "这个分级适用于整段"，不是数据缺失，导出时必须物化；否则 33 行里有
    25 行 tier 为空，而脚本会"绿着错"。

    但"合并段的锚点本身就是空的"是另一回事（Intel 区 B32:B37 原表确实没写
    分级）—— 那是真空白，如实留空，只降级为 NOTE，不算错误。两者必须分开。

    label: 非空时作为报错前缀，供回归测试标注"预期报错"，免得 CI 日志里
           那条故意触发的报错被误读成真失败。
    """
    def e(msg):
        err((label + msg) if label else msg)

    print("\n3b) 显卡品牌：阵营 / 品牌分级必须逐行物化  [⑪]")
    body = tables.get("gpu-brands.csv") or []
    if not body:
        e("gpu-brands.csv 没有数据行")
        return
    h = BRAND_HEADER
    gi, ti = h.index("vendor_group"), h.index("tier")
    bad_v = [i + 2 for i, r in enumerate(body) if not r[gi].strip()]
    blank_t = [r for r in body if not r[ti].strip()]
    if bad_v:
        e("gpu-brands.csv 有 %d 行 vendor_group 为空（第 %s 行）" % (len(bad_v), bad_v[:8]))
    else:
        ok("vendor_group 全部非空（%d 行）" % len(body), verbose)

    # 空 tier 分成两类：合并段锚点本来就空（真空白） vs 该行落在有值的合并段内却没值（bug）
    by_design, not_materialized = [], []
    if blank_t:
        merged, sheet_rows = {}, None
        try:
            import openpyxl
            ws = openpyxl.load_workbook(XLSX)["显卡品牌"]
            for rng in ws.merged_cells.ranges:
                if rng.min_col == 2:                       # 只看 B 列的合并段
                    anchor = norm(ws.cell(row=rng.min_row, column=2).value)
                    for rr in range(rng.min_row, rng.max_row + 1):
                        merged[rr] = anchor
            # CSV 的每一行对应原表哪一行：按 C 列品牌名顺序对齐。
            # 注意排除分区表头行（C='显卡品牌'），它不进 CSV —— 漏排会让整个对齐错位。
            sheet_rows = [r for r in range(3, ws.max_row + 1)
                          if ws.cell(row=r, column=3).value not in (None, "")
                          and norm(ws.cell(row=r, column=3).value) != "显卡品牌"]
        except Exception as ex:
            warn("无法读 xlsx 判断空 tier 的类别（%s），一律按错误处理" % ex)

        if sheet_rows is None:
            not_materialized = list(blank_t)
        elif len(sheet_rows) != len(body):
            e("gpu-brands.csv 有 %d 行，原表品牌数据行有 %d 行，数量对不上 —— "
              "很可能是导出时把某个表头行当成数据、或漏掉了某行" % (len(body), len(sheet_rows)))
            not_materialized = list(blank_t)
        else:
            for r, srow in zip(body, sheet_rows):
                if r[ti].strip():
                    continue
                if merged.get(srow, "") != "":
                    not_materialized.append(r)             # 落在有值的合并段内却没值 → 没物化
                else:
                    by_design.append(r)                    # 该段锚点本来就空 → 真空白

    if not_materialized:
        e("gpu-brands.csv 有 %d/%d 行 tier 为空，且这些行落在 B 列合并段内 —— "
          "说明合并标签没被物化。示例：%s"
          % (len(not_materialized), len(body),
             "、".join("%s(%s)" % (r[2], r[0]) for r in not_materialized[:8])))
    if by_design:
        notes.append("gpu-brands.csv 有 %d 行 tier 为空，均为原表该合并段锚点本身未写分级（真空白，如实留空）：%s"
                     % (len(by_design), "、".join("%s(%s)" % (r[2], r[0]) for r in by_design)))
    if not not_materialized:
        ok("tier 无「该有却没有」的空值（%d 行，合并段已物化%s）"
           % (len(body), "；另有 %d 行原表本就未写分级" % len(by_design) if by_design else ""))


# ------------------------------------------------------------------ ⑧⑩
def check_against_xlsx(tables, verbose):
    print("\n4) 与原工作簿交叉验证：内存聚合值 / 显卡涨幅  [⑧ ⑩]")
    try:
        import openpyxl
    except ImportError:
        warn("没装 openpyxl，跳过与 xlsx 的交叉验证")
        return
    try:
        wb = openpyxl.load_workbook(XLSX, data_only=True)
    except Exception as e:
        warn("读不到 %s：%s" % (XLSX, e))
        return

    # 内存：AVG / MAX / MIN 用 CSV 里的原始单价重算
    body = tables.get("memory-prices.csv") or []
    groups = defaultdict(dict)
    for r in body:
        date, cat, spec = r[0], r[1], r[3]
        if spec in ("AVG", "MAX", "MIN"):
            continue
        groups[(cat, date)][spec] = float(r[4])
    if "内存行情" in wb.sheetnames:
        ws = wb["内存行情"]
        wsf = openpyxl.load_workbook(XLSX, data_only=False)["内存行情"]
        bad, checked, stale = 0, 0, []
        year_of, month_of, cur = {}, {}, None
        from openpyxl.utils import get_column_letter as gl
        for c in range(1, ws.max_column + 1):
            y = ws.cell(row=1, column=c).value
            if isinstance(y, (int, float)) and not isinstance(y, bool):
                cur = int(y)
            year_of[c] = cur
            m = ws.cell(row=2, column=c).value
            if isinstance(m, (int, float)) and not isinstance(m, bool):
                month_of[c] = int(m)
        for r, (cat, spec, kind) in {6: ("内存 DDR4", "AVG", "avg"), 12: ("内存 DDR5", "MAX", "max"),
                                     13: ("内存 DDR5", "MIN", "min"), 16: ("固态 PCIE4.0", "AVG", "avg"),
                                     19: ("固态 PCIE5.0", "AVG", "avg"), 23: ("机械硬盘 HDD", "AVG", "avg")}.items():
            for c, month in month_of.items():
                exp = ws.cell(row=r, column=c).value
                if not isinstance(exp, (int, float)):
                    continue
                key = (cat, "%d-%02d" % (year_of[c], month))
                vals = list(groups.get(key, {}).values())
                if not vals:
                    continue
                got = {"avg": sum(vals) / len(vals), "max": max(vals), "min": min(vals)}[kind]
                checked += 1
                if abs(got - float(exp)) > 0.011:
                    bad += 1
                    stale.append("%s %s %s：按 CSV 原始单价应为 %.3f，xlsx 里是 %.3f（公式 %s）"
                                 % (cat, key[1], kind.upper(), got, float(exp),
                                    wsf.cell(row=r, column=c).value))
        if checked:
            ok("内存/存储聚合值 %d 处重算对照完成（AVG / MAX / MIN）" % checked)
        if stale:
            for s in stale[:6]:
                notes.append("xlsx 缓存值与公式输入不一致（打开重算即可）：" + s)
            if len(stale) > 6:
                notes.append("另有 %d 处同类不一致" % (len(stale) - 6))
        elif not checked:
            warn("内存聚合值没有可校验的交集")

    # 显卡涨幅：把原表公式的基准列与 CSV 的"表内历史最低价"对照
    if "显卡行情" in wb.sheetnames and "gpu-prices.csv" in tables:
        ws = wb["显卡行情"]
        body = tables["gpu-prices.csv"]
        series = defaultdict(dict)
        for r in body:
            if r[6] == "7 月下半月快照":
                continue
            series[r[2]][r[0]] = float(r[4])
        mismatch, checked = 0, 0
        for r in range(3, 42):
            model = ws.cell(row=r, column=3).value
            yform = ws.cell(row=r, column=25).value
            if not model or not isinstance(yform, str):
                continue
            m = re.match(r"=TEXT\(\((\w+)\d+-(\w+)\d+\)/(\w+)\d+", yform)
            if not m:
                continue
            cur_col, base, _ = m.groups()
            from openpyxl.utils import column_index_from_string as ci
            y_cur = ws.cell(row=1, column=ci(cur_col[0])).value
            for c in range(1, ws.max_column + 1):
                if ws.cell(row=1, column=c).value == y_cur:
                    break
            y_now = ws.cell(row=1, column=25).value
            price_cur = ws.cell(row=r, column=ci(cur_col[0])).value
            base_val = ws.cell(row=r, column=ci(base[0])).value
            if not isinstance(price_cur, (int, float)) or not isinstance(base_val, (int, float)):
                continue
            model_name = str(model)
            hist = min(series.get(model_name, {}).values()) if series.get(model_name) else None
            checked += 1
            if hist is not None and abs(hist - float(base_val)) > 0.5:
                mismatch += 1
                notes.append("涨幅基准不一致：%s 的公式用 %s%d=%.0f 作基准，"
                             "而表内历史最低价是 %.0f（来自 CSV）"
                             % (model_name, base[0], r, float(base_val), hist))
        if checked:
            ok("显卡涨幅公式解析 %d 条" % checked)
            if mismatch:
                warn("其中 %d 条的公式基准列不是『表内历史最低价』——"
                     "请确认涨幅口径（详见文末说明）" % mismatch)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args()
    print("=" * 78)
    print("check_csv.py —— data/*.csv 校验")
    print("源工作簿：%s" % os.path.relpath(XLSX, ROOT))
    print("=" * 78)
    tables = check_shape(args.verbose)
    check_dates(tables, args.verbose)
    check_tier_math(tables, args.verbose)
    check_brands(tables, args.verbose)
    check_against_xlsx(tables, args.verbose)

    print()
    print("=" * 78)
    for n in notes[:20]:
        print("  [NOTE] " + n)
    print("错误 %d 条，警告 %d 条" % (len(errors), len(warnings)))
    if errors:
        print("结论：不通过")
        return 1
    print("结论：通过" + ("（有警告，建议看一眼）" if warnings else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main())
